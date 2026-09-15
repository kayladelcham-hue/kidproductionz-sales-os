"""V5F read-only lead input shape validator."""
import argparse,csv,json
from pathlib import Path
from openpyxl import load_workbook
from v5e_execution_adapter import build
REQUIRED={'name','city','state','category'}; ALIASES={'business_name':'name','company':'name','phone_number':'phone','website_url':'website','instagram_url':'instagram'}
def normalize_text(value):
    return '' if value is None else str(value).strip()
def validate(root,campaign,source,sheet=None):
 build(root,campaign); p=Path(source); errors=[];warnings=[]
 if not p.exists(): return {'status':'REJECTED','source_file':str(p),'blocking_errors':['FILE_NOT_FOUND']}
 try:
  if p.suffix.lower()=='.xlsx':
   wb=load_workbook(p,read_only=True,data_only=True); usable=[s for s in wb.sheetnames if any(any(c.value is not None for c in row) for row in wb[s].iter_rows())]
   if not usable: return {'status':'REJECTED','source_file':str(p),'blocking_errors':['EMPTY_WORKBOOK']}
   if sheet:
    if sheet not in usable:return {'status':'REJECTED','source_file':str(p),'blocking_errors':['SHEET_NOT_FOUND:'+sheet]}
   elif len(usable)!=1:return {'status':'REJECTED','source_file':str(p),'blocking_errors':['MULTIPLE_SHEETS_REQUIRE_SELECTION'],'available_sheets':usable}
   ws=wb[sheet or usable[0]]; data=list(ws.values); fields=[str(x or '').strip() for x in data[0]] if data else []; rows=[dict(zip(fields,row)) for row in data[1:]]
  else:
   with p.open(encoding='utf-8-sig',newline='') as f: reader=csv.DictReader(f); fields=reader.fieldnames or []; rows=list(reader)
 except Exception as e:return {'status':'REJECTED','source_file':str(p),'blocking_errors':['READ_ERROR:'+type(e).__name__]}
 norm=[ALIASES.get(normalize_text(x).lower(),normalize_text(x).lower()) for x in fields]; dup=[x for x in set(norm) if norm.count(x)>1]; missing=sorted(REQUIRED-set(norm));
 if dup: errors.append('DUPLICATE_COLUMNS:'+','.join(sorted(dup)))
 if missing: errors.append('MISSING_COLUMNS:'+','.join(missing))
 keys=[tuple(normalize_text(r.get(f)).lower() for f in fields) for r in rows]; duplicates=len(keys)-len(set(keys));
 if duplicates: warnings.append('DUPLICATE_ROWS:'+str(duplicates))
 bad=sum(len(r)!=len(fields) for r in rows); 
 if bad: errors.append('MALFORMED_ROWS:'+str(bad))
 manifest={'campaign_id':campaign,'source_file':str(p.resolve()),'input_schema_version':1,'row_count':len(rows),'normalized_column_map':dict(zip(fields,norm)),'available_channels':sorted(set(x for x in ('phone','email','website','instagram','facebook','tiktok') if x in norm)),'market_coverage':sorted(set((normalize_text(r.get('city'))+', '+normalize_text(r.get('state'))) for r in rows if r.get('city') is not None or r.get('state') is not None)),'category_coverage':sorted(set(normalize_text(r.get('category')) for r in rows if r.get('category') is not None)),'duplicate_summary':{'duplicate_rows':duplicates,'duplicate_columns':dup},'invalid_row_summary':{'malformed_rows':bad},'warnings':warnings,'blocking_errors':errors,'status':'REJECTED' if errors else 'ACCEPTED_FOR_PROCESSING'}
 return manifest
if __name__=='__main__':
 a=argparse.ArgumentParser();a.add_argument('--campaign',required=True);a.add_argument('--input',required=True);a.add_argument('--sheet');a.add_argument('--dry-run',action='store_true');x=a.parse_args();root=Path(__file__).resolve().parents[1];print(json.dumps(validate(root,x.campaign,x.input,x.sheet),indent=2))
