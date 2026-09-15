"""Read source files without moving, modifying, or evaluating them."""
import csv
from pathlib import Path

def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp1252")
        import io
        try:
            dialect = csv.Sniffer().sniff(content[:65536], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(content, newline=""), dialect)
        headers = next(reader, [])
        validate_headers(headers)
        for number, cells in enumerate(reader, 2):
            if not any(str(c).strip() for c in cells):
                continue
            if len(cells) != len(headers):
                raise ValueError(f"CSV row {number}: expected {len(headers)} cells, got {len(cells)}")
            yield "", number, dict(zip(headers, cells))
    elif path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX support requires: pip install -r requirements.txt") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                if not any(v is not None for v in headers):
                    continue
                headers = [str(v).strip() if v is not None else "" for v in headers]
                validate_headers(headers)
                for number, cells in enumerate(rows, 2):
                    if any(c is not None and str(c).strip() for c in cells):
                        yield sheet.title, number, dict(zip(headers, cells))
        finally:
            workbook.close()
    else:
        raise ValueError(f"Unsupported input format: {path.suffix}")

def validate_headers(headers):
    from normalize import key
    names = [key(h) for h in headers]
    if not names or any(not h for h in names) or len(names) != len(set(names)):
        raise ValueError("Input needs a non-empty, unique header for every column.")
