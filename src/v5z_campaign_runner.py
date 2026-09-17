"""V5Z safe, local campaign runner (dry-run only)."""
import argparse, csv, json
from pathlib import Path
from openpyxl import load_workbook
from v5e_execution_adapter import build
from v5f_input_adapter import validate
from scoring import evaluate
from qualification_config import load as load_qualification_config
from outreach_routing import routing
from v5x_queue import build as build_queue
from v5y_queue_artifacts import write as write_queue

def normalize_state(value):
    raw='' if value is None else str(value).strip()
    aliases={'fl':'FL','florida':'FL','ga':'GA','georgia':'GA'}
    return aliases.get(raw.casefold(), raw.upper())

def _rows(path, sheet=None):
    p=Path(path)
    if p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, read_only=True, data_only=True)
        names=[n for n in wb.sheetnames if any(any(c.value is not None for c in row) for row in wb[n].iter_rows())]
        name=sheet or (names[0] if len(names)==1 else None)
        if not name: raise ValueError('SHEET_SELECTION_REQUIRED')
        data=list(wb[name].values); fields=[str(x or '').strip() for x in data[0]]
        return [dict(zip(fields,row)) for row in data[1:]]
    with p.open(encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))

def run(root, campaign, source, sheet=None, dry_run=False):
    if not dry_run: raise ValueError('DRY_RUN_REQUIRED')
    spec=build(root,campaign)
    safety=spec['safety']
    for key in ('network_execution','hubspot_reads','hubspot_writes','outbound_actions'):
        if safety.get(key): raise ValueError('SAFETY_GATE:'+key)
    manifest=validate(root,campaign,source,sheet)
    if manifest.get('status')!='ACCEPTED_FOR_PROCESSING':
        return {'status':'REJECTED','campaign_id':campaign,'input':manifest,'safety':safety}
    cfg=load_qualification_config(root,campaign)
    rcfg=json.loads((root/'config/outreach_config.json').read_text(encoding='utf-8'))
    records=[]
    for i,raw in enumerate(_rows(source,sheet)):
        rec=dict(raw); rec['_source_state']=raw.get('state'); rec['state']=normalize_state(raw.get('state')); rec.update({'website':raw.get('website',''),'phone':raw.get('phone',''),'email':raw.get('email',''),'social':raw.get('instagram',raw.get('social','')),'domain':raw.get('website',''),'other_contact':'','status':'','permanently_closed':'','temporarily_closed':'','rating':raw.get('rating'),'reviews':raw.get('reviews'),'visual':raw.get('visual',''),'visual_evidence':raw.get('visual_evidence',''),'ownership':raw.get('ownership',''),'ownership_evidence':raw.get('ownership_evidence',''),'owner':raw.get('owner','') or raw.get('owner_title','')})
        result=evaluate(rec,cfg); routed=dict(rec); routed.update(result); route,reason,clean=routing(routed,rcfg)
        final=dict(rec); final.update(result); final.update({'route':route,'route_reason':reason,'phone':clean.get('phone',''),'website':clean.get('website',''),'social':clean.get('social','')}); records.append(final)
    queue=build_queue(root,campaign,records)
    outdir=root/'validation_outputs'/'v5_queue'/campaign
    refs=write_queue(queue,outdir)
    run_doc={'run_schema_version':'1.0','campaign_id':campaign,'input':{'source_file':Path(source).name,'row_count':manifest['row_count'],'status':manifest['status']},'configuration_validation_status':spec['validation_status'],'runtime_spec_summary':{'runtime_spec_version':spec['runtime_spec_version']},'qualification_scoring_summary':{'qualified_count':queue['summary']['total_candidates']-queue['summary']['ineligible_count']},'routing_summary':queue.get('summary',{}),'priority_summary':{},'daily_queue_summary':queue['summary'],'artifact_references':{'queue_json':str(refs[0].relative_to(root)),'queue_csv':str(refs[1].relative_to(root)),'queue_markdown':str(refs[2].relative_to(root))},'safety':{'dry_run':True,'network':False,'brave':False,'hubspot_reads':False,'hubspot_writes':False,'outbound':False},'overall_status':'COMPLETED_DRY_RUN'}
    run_dir=root/'validation_outputs'/'v5_runs'/campaign; run_dir.mkdir(parents=True,exist_ok=True); version=1
    run_doc['run_id']=f'{campaign}-v{version}'
    text=json.dumps(run_doc,indent=2,sort_keys=True)+'\n'
    while (run_dir/f'run_v{version}.json').exists() and (run_dir/f'run_v{version}.json').read_text(encoding='utf-8')!=text:
        version+=1; run_doc['run_id']=f'{campaign}-v{version}'; text=json.dumps(run_doc,indent=2,sort_keys=True)+'\n'
    rp=run_dir/f'run_v{version}.json'
    if not rp.exists(): rp.write_text(text,encoding='utf-8')
    try:
        from app.api.database import persist_run_bundle
        persist_run_bundle(run_doc, records, queue)
    except Exception as exc:
        run_doc['persistence_warning']='Database persistence failed: '+type(exc).__name__
    return run_doc

if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--campaign',required=True); ap.add_argument('--input',required=True); ap.add_argument('--sheet'); ap.add_argument('--dry-run',action='store_true'); a=ap.parse_args()
    try: print(json.dumps(run(Path(__file__).resolve().parents[1],a.campaign,a.input,a.sheet,a.dry_run),indent=2))
    except Exception as e: print(json.dumps({'status':'FAILED_SAFE','error':str(e)})); raise SystemExit(1)
