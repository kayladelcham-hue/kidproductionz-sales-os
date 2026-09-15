import unittest,csv,tempfile
from pathlib import Path
from v5f_input_adapter import validate
from openpyxl import Workbook
class V5F(unittest.TestCase):
 def test_valid_and_missing(self):
  root=Path(__file__).parents[1]
  with tempfile.NamedTemporaryFile('w',suffix='.csv',delete=False,newline='') as f:
   w=csv.DictWriter(f,fieldnames=['name','city','state','category','phone']);w.writeheader();w.writerow({'name':'A','city':'Orlando','state':'FL','category':'hair salon','phone':'1'})
   p=f.name
  self.assertEqual(validate(root,'orlando_beauty',p)['status'],'ACCEPTED_FOR_PROCESSING')
  with tempfile.NamedTemporaryFile('w',suffix='.csv',delete=False) as f:f.write('name,city\nA,Orlando');q=f.name
  self.assertEqual(validate(root,'orlando_beauty',q)['status'],'REJECTED')
 def test_xlsx_and_sheet_rules(self):
  root=Path(__file__).parents[1]; p=Path(tempfile.mktemp(suffix='.xlsx')); w=Workbook(); ws=w.active;ws.title='Businesses';ws.append(['business_name','city','state','category','email','website']);ws.append(['A','Orlando','FL','hair salon','a@x.com','https://x.com']);w.save(p)
  r=validate(root,'orlando_beauty',p);self.assertEqual(r['status'],'ACCEPTED_FOR_PROCESSING');self.assertIn('name',r['normalized_column_map'].values())
  p2=Path(tempfile.mktemp(suffix='.xlsx'));w=Workbook();w.active.append(['name','city','state','category']);w.active.append(['A','Orlando','FL','hair salon']);s=w.create_sheet('Other');s.append(['name','city','state','category']);s.append(['B','Orlando','FL','hair salon']);w.save(p2);self.assertEqual(validate(root,'orlando_beauty',p2)['status'],'REJECTED');self.assertEqual(validate(root,'orlando_beauty',p2,'Other')['status'],'ACCEPTED_FOR_PROCESSING')
 def test_channels_duplicates_and_empty(self):
  root=Path(__file__).parents[1];p=Path(tempfile.mktemp(suffix='.csv'));p.write_text('name,city,state,category,phone,email,website,instagram\nA,Atlanta,GA,hair salon,1,a@x.com,x.com,@a\nA,Atlanta,GA,hair salon,1,a@x.com,x.com,@a\n');r=validate(root,'atlanta_beauty',p);self.assertIn('phone',r['available_channels']);self.assertIn('email',r['available_channels']);self.assertGreater(r['duplicate_summary']['duplicate_rows'],0)
if __name__=='__main__':unittest.main()
