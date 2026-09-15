import unittest,csv,tempfile
from pathlib import Path
from openpyxl import Workbook
from v5f_input_adapter import validate
class V5I(unittest.TestCase):
 def setUp(self): self.root=Path(__file__).parents[1]
 def csv(self,text):
  p=Path(tempfile.mktemp(suffix='.csv'));p.write_text(text,encoding='utf-8');return p
 def test_empty_malformed_duplicates_and_unknowns(self):
  self.assertEqual(validate(self.root,'orlando_beauty',self.csv(''))['status'],'REJECTED')
  p=self.csv('name,name,city,state,category\nA,A,Orlando,FL,unknown\n');self.assertEqual(validate(self.root,'orlando_beauty',p)['status'],'REJECTED')
  p=self.csv('name,city,state,category\nA,Orlando,FL,hair salon\nA,Orlando,FL,hair salon\n');self.assertEqual(validate(self.root,'orlando_beauty',p)['duplicate_summary']['duplicate_rows'],1)
 def test_xlsx_empty_and_multiple(self):
  p=Path(tempfile.mktemp(suffix='.xlsx'));Workbook().save(p);self.assertEqual(validate(self.root,'orlando_beauty',p)['status'],'REJECTED')
  p=Path(tempfile.mktemp(suffix='.xlsx'));w=Workbook();w.active.append(['name','city','state','category']);w.active.append(['A','Orlando','FL','hair salon']);s=w.create_sheet('Two');s.append(['name','city','state','category']);s.append(['B','Orlando','FL','hair salon']);w.save(p);self.assertIn('MULTIPLE_SHEETS_REQUIRE_SELECTION',validate(self.root,'orlando_beauty',p)['blocking_errors'])
 def test_snapshot_and_equivalence(self):
  p=self.csv('name,city,state,category,phone\nA,Orlando,FL,hair salon,1\n');a=validate(self.root,'orlando_beauty',p);b=validate(self.root,'orlando_beauty',p);self.assertEqual({k:v for k,v in a.items() if k!='source_file'},{k:v for k,v in b.items() if k!='source_file'})
 def test_market_and_missing_fields(self):
  p=self.csv('name,city,state,category\nA,Atlanta,GA,hair salon\n');self.assertEqual(validate(self.root,'orlando_beauty',p)['status'],'ACCEPTED_FOR_PROCESSING')
  p=self.csv('name,city,state\nA,Orlando,FL\n');self.assertEqual(validate(self.root,'orlando_beauty',p)['status'],'REJECTED')
if __name__=='__main__':unittest.main()
